"""
Gera o .pptx da apresentacao para a Trase a partir do outline aprovado
(presentation/outline_pptx_trase.md) e do template visual da MMA.

NOTA: o plano original (presentation/outline_pptx_trase.md + plano de execucao)
previa um script R com officer/flextable para a Etapa 2, pensado para rodar no
RStudio local do usuario (que tem R instalado). Este ambiente de execucao remota
nao tem R disponivel, entao esta primeira geracao foi feita em Python
(python-pptx + matplotlib + pandas), reproduzindo a mesma paleta de cores e a
mesma logica de calculo do R/deforestation_attribution.R. Uma versao R
equivalente pode ser escrita depois, se o usuario quiser rodar/regenerar
localmente com o restante do pipeline em R.

Uso:
    python3 presentation/build_pptx.py \
        --template <caminho para template_mma_dpcd.pptx> \
        --data <caminho para serie_completa_2013_2023.csv> \
        --output <caminho de saida .pptx>
"""

import argparse
import copy
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

# ----- Constantes compartilhadas com R/deforestation_attribution.R -----------
BIOMA_ORDER = ["AMAZONIA", "CERRADO", "MATA ATLANTICA", "CAATINGA", "PAMPA", "PANTANAL"]
CORES = {"soja": "#E69F00", "pasto": "#009E73", "outros": "#999999"}
PERIODOS = [2013, 2018, 2023]
JANELAS = {2013: "2009-2013", 2018: "2014-2018", 2023: "2019-2023"}

EMU_PER_IN = 914400


def emu(x):
    return Emu(int(round(x * EMU_PER_IN)))


# ----- Carregar dados ---------------------------------------------------------
def load_data(csv_path):
    df = pd.read_csv(csv_path)
    df = df[df["pasture_year"].isin(PERIODOS)].copy()
    df["janela_desmatamento"] = df["pasture_year"].map(JANELAS)
    return df


def fmt_ha(x):
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def fmt_pct(x):
    return f"{x * 100:,.2f}".replace(".", ",") + "%"


# ----- Graficos (matplotlib, mesma paleta do ggplot em R) --------------------
def make_brasil_charts(df, out_dir):
    br = df[df["bioma"] == "BRASIL"].sort_values("pasture_year")
    janelas = br["janela_desmatamento"].tolist()

    # % empilhado
    fig, ax = plt.subplots(figsize=(8, 5.5), dpi=200)
    soja = br["pct_soy"].values
    pasto = br["pct_pasture"].values
    outros = br["pct_outros"].values
    ax.bar(janelas, soja, color=CORES["soja"], label="soja")
    ax.bar(janelas, pasto, bottom=soja, color=CORES["pasto"], label="pasto")
    ax.bar(janelas, outros, bottom=soja + pasto, color=CORES["outros"], label="outros")
    ax.set_ylim(0, 1)
    ax.set_yticks([0, 0.25, 0.5, 0.75, 1.0])
    ax.set_yticklabels(["0%", "25%", "50%", "75%", "100%"])
    ax.set_title("Brasil: % do desmatamento atribuido a soja e pasto", fontweight="bold", loc="left")
    ax.text(0, 1.06, "Tres janelas de 5 anos nao sobrepostas", transform=ax.transAxes, fontsize=9, color="dimgray")
    ax.set_xlabel("Janela de desmatamento")
    ax.set_ylabel("% do desmatamento territorial da janela")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.15), ncol=3, frameon=False)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    p1 = out_dir / "brasil_pct_3_periodos.png"
    fig.savefig(p1)
    plt.close(fig)

    # hectares empilhado (milhoes)
    fig, ax = plt.subplots(figsize=(8, 5.5), dpi=200)
    soja_ha = br["soy_deforestation_ha"].values / 1e6
    pasto_ha = br["pasture_deforestation_ha"].values / 1e6
    total_ha = br["desmatamento_territorial_total_ha"].values / 1e6
    outros_ha = (total_ha - soja_ha - pasto_ha).clip(min=0)
    ax.bar(janelas, soja_ha, color=CORES["soja"], label="soja")
    ax.bar(janelas, pasto_ha, bottom=soja_ha, color=CORES["pasto"], label="pasto")
    ax.bar(janelas, outros_ha, bottom=soja_ha + pasto_ha, color=CORES["outros"], label="outros")
    ax.set_title("Brasil: desmatamento atribuido a soja e pasto (hectares)", fontweight="bold", loc="left")
    ax.text(0, 1.06, "Tres janelas de 5 anos nao sobrepostas", transform=ax.transAxes, fontsize=9, color="dimgray")
    ax.set_xlabel("Janela de desmatamento")
    ax.set_ylabel("Milhoes de hectares")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.15), ncol=3, frameon=False)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    p2 = out_dir / "brasil_hectares_3_periodos.png"
    fig.savefig(p2)
    plt.close(fig)
    return p1, p2


def make_facet_chart(df, out_dir):
    biomas_df = df[df["bioma"].isin(BIOMA_ORDER)].copy()
    fig, axes = plt.subplots(1, 3, figsize=(13, 5.5), dpi=200, sharey=True)
    for ax, px in zip(axes, PERIODOS):
        janela = JANELAS[px]
        sub = biomas_df[biomas_df["pasture_year"] == px].set_index("bioma").reindex(BIOMA_ORDER)
        soja = sub["pct_soy"].values
        pasto = sub["pct_pasture"].values
        outros = sub["pct_outros"].values
        ax.bar(BIOMA_ORDER, soja, color=CORES["soja"], label="soja")
        ax.bar(BIOMA_ORDER, pasto, bottom=soja, color=CORES["pasto"], label="pasto")
        ax.bar(BIOMA_ORDER, outros, bottom=soja + pasto, color=CORES["outros"], label="outros")
        ax.set_title(janela, fontsize=11)
        ax.set_ylim(0, 1)
        ax.set_yticks([0, 0.25, 0.5, 0.75, 1.0])
        ax.set_yticklabels(["0%", "25%", "50%", "75%", "100%"])
        ax.tick_params(axis="x", rotation=30)
        for label in ax.get_xticklabels():
            label.set_ha("right")
        ax.spines[["top", "right"]].set_visible(False)
    axes[0].set_ylabel("% do desmatamento territorial da janela")
    fig.suptitle("% do desmatamento atribuido a soja e pasto, por bioma", fontweight="bold", x=0.01, ha="left")
    fig.text(0.01, 0.93, "Tres janelas de 5 anos nao sobrepostas", fontsize=9, color="dimgray")
    handles, labels = axes[0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="lower center", ncol=3, frameon=False, bbox_to_anchor=(0.5, -0.02))
    fig.tight_layout(rect=[0, 0.04, 1, 0.90])
    p = out_dir / "comparacao_3_periodos.png"
    fig.savefig(p)
    plt.close(fig)
    return p


# ----- Helpers de slide --------------------------------------------------------
def get_layout(prs, name):
    for layout in prs.slide_masters[0].slide_layouts:
        if layout.name == name:
            return layout
    raise KeyError(f"Layout '{name}' nao encontrado no template")


def clear_existing_slides(prs):
    """Remove todas as slides do template (conteudo + relacionamento), mantendo
    apenas masters/layouts. Sem dropar o relationship, a parte antiga do slide
    continua no pacote e colide com o partname reatribuido as slides novas."""
    xml_slides = prs.slides._sldIdLst
    for sld in list(xml_slides):
        rId = sld.get(qn("r:id"))
        prs.part.drop_rel(rId)
        xml_slides.remove(sld)


def strip_placeholder(slide, idx):
    """Remove um placeholder vazio do slide (para substituir por shape livre)."""
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx == idx:
            ph._element.getparent().remove(ph._element)
            return


def set_title(slide, text):
    slide.shapes.title.text = text


def add_bullets(slide, box, lines, font_size=16, bullet="•  "):
    tb = slide.shapes.add_textbox(*box)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"{bullet}{line}" if bullet else line
        p.font.size = Pt(font_size)
        p.space_after = Pt(8)
    return tb


def add_table(slide, box, header, rows):
    left, top, width, height = box
    n_rows = len(rows) + 1
    n_cols = len(header)
    gtable = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = gtable.table
    for j, h in enumerate(header):
        cell = table.cell(0, j)
        cell.text = h
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].font.size = Pt(12)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(val)
            cell.text_frame.paragraphs[0].font.size = Pt(11)
    return gtable


# ----- Montagem do pptx --------------------------------------------------------
def build(template_path, data_path, output_path, assets_dir):
    df = load_data(data_path)
    br = df[df["bioma"] == "BRASIL"].sort_values("pasture_year")
    tabela_rows = [
        [
            row.janela_desmatamento,
            fmt_ha(row.desmatamento_territorial_total_ha),
            fmt_ha(row.soy_deforestation_ha),
            fmt_pct(row.pct_soy),
            fmt_ha(row.pasture_deforestation_ha),
            fmt_pct(row.pct_pasture),
            fmt_pct(row.pct_outros),
        ]
        for row in br.itertuples()
    ]
    header = ["Janela", "Desmat. territorial (ha)", "Soja (ha)", "% Soja", "Pasto (ha)", "% Pasto", "% Outros"]

    assets_dir.mkdir(parents=True, exist_ok=True)
    img_br_pct, img_br_ha = make_brasil_charts(df, assets_dir)
    img_facet = make_facet_chart(df, assets_dir)

    prs = Presentation(template_path)
    clear_existing_slides(prs)

    L_TITLE = get_layout(prs, "Slide de Título")
    L_CONTENT = get_layout(prs, "Título e Conteúdo")
    L_TWOCONTENT = get_layout(prs, "Duas Partes de Conteúdo")
    L_TITLEONLY = get_layout(prs, "Somente Título")

    body_box = (emu(0.92), emu(2.0), emu(11.5), emu(4.76))
    two_box_1 = (emu(0.92), emu(2.0), emu(5.67), emu(4.76))
    two_box_2 = (emu(6.75), emu(2.0), emu(5.67), emu(4.76))
    wide_img_box = (emu(0.92), emu(1.6), emu(11.5), emu(5.4))

    # Slide 1 - Capa
    s = prs.slides.add_slide(L_TITLE)
    s.placeholders[0].text = "Atribuição do desmatamento territorial a soja e pasto"
    s.placeholders[1].text = (
        "Pré-leitura para a reunião Trase × DPCD/MMA sobre PRODES e commodities\n"
        "DPCD/CGIE — Ministério do Meio Ambiente"
    )

    # Slide 2 - Contexto e pergunta de pesquisa
    s = prs.slides.add_slide(L_CONTENT)
    set_title(s, "Contexto e pergunta de pesquisa")
    strip_placeholder(s, 1)
    add_bullets(s, body_box, [
        "Em breve a Trase apresenta ao DPCD/MMA a relação entre o PRODES e commodities. "
        "Numa conversa preliminar, ficou levantada a preocupação de que a categoria "
        "“não-associado” (outros) aparece muito grande.",
        "Pergunta de pesquisa: qual % do desmatamento territorial, por bioma e para o "
        "Brasil, é atribuível a soja e a pasto — aplicando estritamente a metodologia de "
        "janelas de 5 anos definida pela própria Trase — e o que resta como resíduo "
        "(“outros”)?",
        "Escopo: 6 biomas (Amazônia, Cerrado, Mata Atlântica, Caatinga, Pampa, Pantanal) "
        "+ Brasil; 3 janelas de 5 anos não sobrepostas: 2009-2013, 2014-2018, 2019-2023.",
    ], font_size=17)

    # Slide 3 - Fontes de dados exatas
    s = prs.slides.add_slide(L_CONTENT)
    set_title(s, "Fontes de dados exatas")
    strip_placeholder(s, 1)
    add_bullets(s, (emu(0.92), emu(2.0), emu(11.5), emu(0.6)),
                ["Fonte: Trase SEI-PCS Brazil, beef v2.2 / soy v2.6 — granularidade por bioma (region_trase_id), 6 biomas."],
                font_size=15, bullet="")
    add_table(s, (emu(0.92), emu(2.8), emu(11.5), emu(2.2)),
              ["Arquivo", "Cobertura", "Coluna de valor"],
              [
                  ["territorial_deforestation_biome.csv", "2008–2024, anual", "deforestation_hectares"],
                  ["pasture_deforestation_5_year_total_biome.csv", "2010–2023", "pasture_deforestation_hectares"],
                  ["soy_deforestation_5_year_total_biome.csv", "2014–2024", "soy_deforestation_hectares"],
              ])

    # Slide 4 - Método: definição oficial e pareamento
    s = prs.slides.add_slide(L_CONTENT)
    set_title(s, "Método: definição oficial e pareamento")
    strip_placeholder(s, 1)
    add_bullets(s, body_box, [
        "Conforme documentação SEI-PCS Brazil beef v2.2 / soy v2.6:",
        "pasture_deforestation(X) = pasto mapeado em X sobre desmatamento territorial na "
        "janela [X-4, X] (5 anos, inclui X).",
        "soy_deforestation(X) = soja mapeada em X sobre desmatamento territorial na "
        "janela [X-5, X-1] (5 anos, defasagem de 1 ano, não inclui X).",
        "Pareamento: para a mesma janela de desmatamento, casa-se pasture(X) com "
        "soy(X+1) — ambos cobrem [X-4, X].",
        "Só 3 janelas não sobrepostas: “5 year totals” de anos vizinhos "
        "compartilham até 4 anos de desmatamento territorial e não podem ser somados "
        "sem dupla contagem — a série completa ano a ano fica fora de escopo aqui.",
    ], font_size=16)

    # Slide 5 - Formulas exatas aplicadas
    s = prs.slides.add_slide(L_CONTENT)
    set_title(s, "Fórmulas exatas aplicadas")
    strip_placeholder(s, 1)
    tb = s.shapes.add_textbox(*body_box)
    tf = tb.text_frame
    tf.word_wrap = True
    formula_lines = [
        "pct_soy         = soy_deforestation_ha(X+1) / desmatamento_territorial_total_ha[X-4,X]",
        "pct_pasture     = pasture_deforestation_ha(X) / desmatamento_territorial_total_ha[X-4,X]",
        "pct_outros      = 1 - pct_soy - pct_pasture          (sem clamp; pode ser negativo)",
        "overlap_warning = (pct_soy + pct_pasture) > 1",
    ]
    for i, line in enumerate(formula_lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Consolas"
        p.font.size = Pt(15)
        p.space_after = Pt(6)
    p = tf.add_paragraph()
    p.text = (
        "Para o Brasil: somam-se primeiro os hectares dos 6 biomas, e só então "
        "recalculam-se os 3 percentuais a partir dos totais somados — nunca se somam "
        "percentuais de bioma diretamente."
    )
    p.font.size = Pt(15)
    p.space_before = Pt(16)

    # Slide 6 - Resultados: Brasil agregado
    s = prs.slides.add_slide(L_TWOCONTENT)
    set_title(s, "Resultados: Brasil agregado")
    strip_placeholder(s, 1)
    strip_placeholder(s, 2)
    s.shapes.add_picture(str(img_br_pct), *two_box_1)
    s.shapes.add_picture(str(img_br_ha), *two_box_2)

    # Slide 7 - Resultados: por bioma, 3 periodos
    s = prs.slides.add_slide(L_TITLEONLY)
    set_title(s, "Resultados: por bioma, 3 períodos")
    s.shapes.add_picture(str(img_facet), *wide_img_box)

    # Slide 8 - Leitura dos resultados
    s = prs.slides.add_slide(L_CONTENT)
    set_title(s, "Leitura dos resultados")
    strip_placeholder(s, 1)
    add_bullets(s, body_box, [
        "No Brasil, soja + pasto explicam uma fatia crescente do desmatamento "
        "territorial: 47,16% → 54,16% → 57,82% (2009-13 → 2014-18 → 2019-23). "
        "“Outros” caiu de 52,84% para 42,18%, mas ainda é quase metade do "
        "desmatamento no período mais recente.",
        "O desmatamento territorial total cresceu ~40,8% entre o primeiro e o último "
        "período (10,09 → 14,21 milhões de ha).",
        "Por bioma em 2019-2023: Amazônia (78,18%) e Pantanal (76,65%) dominados por "
        "pasto; Pampa é o único bioma soja-dominante (40,10% soja vs. 0,34% pasto); "
        "Mata Atlântica tem o maior resíduo “outros” (74,70%), seguida de "
        "Cerrado (59,80%) e Pampa (59,56%).",
        "Nenhum dos 21 casos analisados (3 períodos × 6 biomas × Brasil) disparou "
        "overlap_warning.",
    ], font_size=16)

    # Slide 9 - Tabela-resumo Brasil, 3 periodos
    s = prs.slides.add_slide(L_CONTENT)
    set_title(s, "Tabela-resumo (Brasil, 3 períodos)")
    strip_placeholder(s, 1)
    add_table(s, (emu(0.92), emu(2.3), emu(11.5), emu(2.2)), header, tabela_rows)

    # Slide 10 - Limitacoes, notas tecnicas e perguntas
    s = prs.slides.add_slide(L_CONTENT)
    set_title(s, "Limitações, notas técnicas e perguntas para a Trase")
    strip_placeholder(s, 1)
    add_bullets(s, (emu(0.92), emu(1.95), emu(11.5), emu(2.35)), [
        "pct_outros não tem clamp e pode ficar negativo em caso de sobreposição "
        "espacial soja/pasto (overlap_warning) — não ocorreu nos períodos aqui.",
        "Assimetria: no gráfico em hectares, “outros” é truncado em zero; nas "
        "tabelas/gráfico percentual não há essa correção.",
        "Cobertura temporal desigual dos 3 CSVs de origem limita o intervalo de "
        "janelas analisável.",
    ], font_size=14)
    add_bullets(s, (emu(0.92), emu(4.4), emu(11.5), emu(2.3)), [
        "Nosso entendimento do pareamento pasture(X)+soy(X+1) e das janelas não "
        "sobrepostas está correto?",
        "Como a Trase recomenda decompor/explicar a categoria “outros” ao "
        "DPCD?",
        "Existe orientação oficial da Trase sobre a sobreposição espacial potencial "
        "soja×pasto?",
        "Há previsão de estender a cobertura temporal de pasture/soy?",
    ], font_size=14)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return output_path


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True)
    parser.add_argument("--data", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--assets", default="presentation/assets")
    args = parser.parse_args()

    out = build(Path(args.template), Path(args.data), Path(args.output), Path(args.assets))
    print(f"PPTX gerado em: {out}")
